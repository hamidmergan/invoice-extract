"""Write rows to Excel, CSV or JSON."""

import csv
import datetime
import json
import re
from dataclasses import asdict
from decimal import Decimal
from pathlib import Path
from typing import Any, List, Optional, Set, Tuple, Union

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .columns import STANDARD_COLUMNS, Column
from .errors import ConfigError
from .results import Row, _value

FORMATS = ("xlsx", "csv", "json")
DETAILS_HEADER = ["File", "Column", "Value", "Raw text", "Source", "Page", "Edited", "Checks"]
# Spreadsheet apps run cells that start with these as formulas.
DANGEROUS = ("=", "+", "-", "@", "\t", "\r")
MAX_WIDTH = 50
SCALARS = (str, int, float, Decimal, datetime.date)


def _dangerous(value: Any) -> bool:
    return isinstance(value, str) and value.startswith(DANGEROUS)


def _cell_value(row: Row, column: Column) -> Any:
    if column.key == "file":
        return Path(row.path).name
    if column.key == "status":
        return row.status
    if row.error:
        return None
    cell = row.cells.get(column.key)
    return None if cell is None else cell.value


def _checks(checks) -> str:
    return "; ".join(f"{c.level}: {c.message}" for c in checks)


def _details(rows: List[Row], columns: List[Column]) -> List[List[Any]]:
    labels = {c.key: c.label for c in columns}
    out = []
    for row in rows:
        name = Path(row.path).name
        if row.error:
            out.append([name, "(error)", None, "", "", None, False, f"error: {row.error}"])
            continue
        for key, cell in row.cells.items():
            if key not in labels:
                continue
            out.append([name, labels.get(key, key), cell.value, cell.raw, cell.source, cell.page, cell.edited,
                        _checks(cell.checks)])
        if row.checks:
            out.append([name, "(invoice)", None, "", "", None, False, _checks(row.checks)])
    return out


def _scalar(value: Any) -> Any:
    """Lists and objects from a model answer become JSON text, so every writer can store them."""
    if value is None or isinstance(value, SCALARS):
        return value
    return json.dumps(value, ensure_ascii=False, default=str)


def _put(ws, row: int, col: int, value: Any):
    value = _scalar(value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)  # control characters from PDF text are not allowed in xlsx
    if isinstance(value, Decimal):
        cell = ws.cell(row, col, float(value))
        cell.number_format = "#,##0.00"
    elif isinstance(value, datetime.date):
        cell = ws.cell(row, col, value)
        cell.number_format = "yyyy-mm-dd"
    else:
        cell = ws.cell(row, col, value)
        if _dangerous(value):
            cell.data_type = "s"
            cell.quotePrefix = True
    return cell


def _sheet(ws, header: List[str], records: List[List[Any]]) -> None:
    widths = [len(h) for h in header]
    for j, label in enumerate(header, start=1):
        _put(ws, 1, j, label).font = Font(bold=True)
    for i, record in enumerate(records, start=2):
        for j, value in enumerate(record, start=1):
            _put(ws, i, j, value)
            if value is not None:
                widths[j - 1] = max(widths[j - 1], len(str(_value(value))))
    for j, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(width + 2, MAX_WIDTH)
    ws.freeze_panes = "A2"


def _plain(columns: List[Column]) -> List[Column]:
    return [c for c in columns if c.kind != "table"]


def _tables(columns: List[Column]) -> List[Column]:
    return [c for c in columns if c.kind == "table"]


def _table(rows: List[Row], column: Column) -> Tuple[List[str], List[List[Any]]]:
    """Header and records of one table field: one record per table row, linked to its invoice."""
    cells = [(row, row.cells.get(column.key)) for row in rows if not row.error]
    names = list(column.table_columns)
    for _, cell in cells:
        if cell is not None:
            names.extend(n for n in cell.columns if n not in names)
    numbered = any(row.cells.get("invoice_no") is not None and row.cells["invoice_no"].value not in (None, "")
                   for row, _ in cells)
    header = ["File"] + (["Invoice No."] if numbered else []) + names
    records = []
    for row, cell in cells:
        if cell is None or not isinstance(cell.value, list):
            continue
        number = row.cells.get("invoice_no")
        lead = [Path(row.path).name] + ([number.value if number is not None else None] if numbered else [])
        for line in cell.value:
            line = line if isinstance(line, dict) else {}
            records.append(lead + [line.get(n) for n in names])
    return header, records


def _sheet_title(label: str, used: Set[str]) -> str:
    """A valid, unused Excel sheet name: no []:*?/\\, at most 31 characters."""
    base = ILLEGAL_CHARACTERS_RE.sub("", re.sub(r"[\[\]:*?/\\]", "", label)).strip().strip("'").strip()
    if not base or base.lower() == "history":  # Excel reserves "History"
        base = (base + " table").strip() or "Table"
    base = base[:31]
    title, n = base, 2
    while title.lower() in {u.lower() for u in used}:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def table_csv_path(path: Union[str, Path], column: Column) -> Path:
    """Where the CSV export writes a table field: next to the main file."""
    path = Path(path)
    return path.with_name(f"{path.stem}-{column.key}.csv")


def _xlsx(rows: List[Row], path: Path, columns: List[Column]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    plain = _plain(columns)
    _sheet(ws, [c.label for c in plain], [[_cell_value(r, c) for c in plain] for r in rows])
    _sheet(wb.create_sheet("Details"), DETAILS_HEADER, _details(rows, plain))
    used = {"Invoices", "Details"}
    for column in _tables(columns):
        header, records = _table(rows, column)
        _sheet(wb.create_sheet(_sheet_title(column.label, used)), header, records)
    wb.save(path)


def _csv_text(value: Any) -> str:
    value = _scalar(value)
    if value is None:
        return ""
    if _dangerous(value):
        return "'" + value
    return str(_value(value))


def _write_csv(path: Path, header: List[str], records: List[List[Any]]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([_csv_text(h) for h in header])
        for record in records:
            writer.writerow([_csv_text(v) for v in record])


def _csv(rows: List[Row], path: Path, columns: List[Column]) -> None:
    plain = _plain(columns)
    _write_csv(path, [c.label for c in plain], [[_cell_value(r, c) for c in plain] for r in rows])
    for column in _tables(columns):
        _write_csv(table_csv_path(path, column), *_table(rows, column))


def _json(rows: List[Row], path: Path, columns: List[Column]) -> None:
    keys = {c.key for c in columns}
    records = []
    for row in rows:
        record = row.to_dict()
        record["cells"] = {k: v for k, v in record["cells"].items() if k in keys}
        records.append(record)
    data = {"columns": [asdict(c) for c in columns], "rows": records}
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


WRITERS = {"xlsx": _xlsx, "csv": _csv, "json": _json}


def export(rows: List[Row], path: Union[str, Path], columns: Optional[List[Column]] = None,
           format: Optional[str] = None) -> Path:
    """Write `rows` to `path`. The format comes from the suffix unless `format` is given."""
    path = Path(path)
    fmt = (format or path.suffix).lower().lstrip(".")
    if fmt not in WRITERS:
        raise ConfigError(f"Unknown export format {fmt!r}; use one of: {', '.join(FORMATS)}")
    WRITERS[fmt](list(rows), path, list(columns or STANDARD_COLUMNS))
    return path
