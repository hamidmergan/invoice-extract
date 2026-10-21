import io
import json
import threading
import time
import zipfile

import openpyxl
import pytest
from pdfgen import make_pdf, write_samples

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402

from invoice_extract.config import resolve  # noqa: E402
from invoice_extract.web.app import create_app  # noqa: E402

PNG = b"\x89PNG\r\n\x1a\n"


def client_for(settings):
    return TestClient(create_app(settings))


def upload(client, *paths):
    files = [("files", (p.name, p.read_bytes(), "application/octet-stream")) for p in paths]
    return client.post("/api/files", files=files)


def wait_done(client, timeout=10.0):
    deadline = time.time() + timeout
    while time.time() < deadline:
        state = client.get("/api/state").json()
        if state["files"] and all(f["state"] == "done" for f in state["files"]):
            return state
        time.sleep(0.05)
    raise AssertionError("files did not finish in time")


def wait_for(predicate, timeout=10.0):
    deadline = time.time() + timeout
    while time.time() < deadline:
        if predicate():
            return
        time.sleep(0.02)
    raise AssertionError("condition not reached in time")


def llm(server):
    return resolve(base_url=server.base_url, model="m", max_retries=0)


def prompt(body):
    content = body["messages"][0]["content"]
    return content if isinstance(content, str) else ""


def customer_pdfs(tmp_path):
    a, b = tmp_path / "a.pdf", tmp_path / "b.pdf"
    a.write_bytes(make_pdf([[(50, 780, "ACME GmbH Invoice RE-1 Total 119,00 EUR"), (50, 700, "Kd.-Nr. 4711")]]))
    b.write_bytes(make_pdf([[(50, 780, "Beta Ltd Invoice B-7 Total 50,00 EUR"), (50, 700, "Customer number 4712")]]))
    return a, b


def discovery_reply(gate=None):
    def reply(body):
        if prompt(body).startswith("You get labels"):
            return json.dumps({"mapping": {"Kd.-Nr.": "Customer No.", "Customer number": "Customer No."}})
        if prompt(body).startswith("You extract values"):
            return json.dumps({"order_no": {"value": "PO-1", "source_text": "PO-1", "page": 1}})
        if gate is not None:
            gate.wait(5)
        text = body["messages"][-1]["content"]
        if "Kd.-Nr. 4711" in text:
            item = {"label": "Kd.-Nr.", "kind": "value", "value": "4711", "source_text": "Kd.-Nr. 4711", "page": 1}
        else:
            item = {"label": "Customer number", "kind": "value", "value": "4712",
                    "source_text": "Customer number 4712", "page": 1}
        return json.dumps({"items": [item]})
    return reply


def test_upload_xml_and_state(tmp_path, server):
    xml = write_samples(tmp_path)[2]
    with client_for(resolve()) as client:
        res = upload(client, xml)
        assert res.status_code == 200
        assert res.json()["files"][0]["name"] == "xrechnung.xml"
        state = wait_done(client)
    entry = state["files"][0]
    assert state["mode"] == "discover"
    assert entry["kind"] == "xml" and entry["status"] == "ok"
    assert entry["row"]["cells"]["invoice_no"]["value"] == "RE-2026-001"
    fields = {f["key"]: f for f in state["fields"]}
    assert {"invoice_no", "gross", "line_items"} <= set(fields)
    assert fields["gross"]["count"] == 1 and fields["gross"]["selected"] is True
    assert fields["line_items"]["kind"] == "table"
    assert state["summary"] == {"files": 1, "fields": len(fields), "to_check": 0}
    assert server.requests == []


def test_missing_model_is_row_error(tmp_path, server):
    pdf = write_samples(tmp_path)[0]
    with client_for(resolve()) as client:
        assert upload(client, pdf).status_code == 200
        state = wait_done(client)
    entry = state["files"][0]
    assert entry["status"] == "error"
    assert "--model" in entry["row"]["error"]
    assert state["summary"]["to_check"] == 1


def test_page_png(tmp_path, server):
    pdf = write_samples(tmp_path)[0]
    with client_for(resolve()) as client:
        file_id = upload(client, pdf).json()["files"][0]["id"]
        assert client.get(f"/api/files/{file_id}/pages").json() == {"count": 1}
        png = client.get(f"/api/files/{file_id}/pages/1.png")
        assert png.status_code == 200
        assert png.headers["content-type"] == "image/png"
        assert png.content.startswith(PNG)
        assert client.get(f"/api/files/{file_id}/pages/2.png").status_code == 404


def test_discovery_state_and_counts(tmp_path, server):
    gate = threading.Event()
    server.reply = discovery_reply(gate)
    with client_for(llm(server)) as client:
        upload(client, *customer_pdfs(tmp_path))
        wait_for(lambda: len(server.requests) > 0)
        states = [f["state"] for f in client.get("/api/state").json()["files"]]
        assert "done" not in states
        gate.set()
        state = wait_done(client)
    fields = {f["key"]: f for f in state["fields"]}
    assert set(fields) == {"customer_no"}
    assert fields["customer_no"]["count"] == 2 and fields["customer_no"]["selected"] is True
    assert [f["row"]["cells"]["customer_no"]["value"] for f in state["files"]] == ["4711", "4712"]


def test_selection_drives_export(tmp_path, server):
    xml = write_samples(tmp_path)[2]
    with client_for(resolve()) as client:
        upload(client, xml)
        wait_done(client)
        assert client.put("/api/selection", json={"keys": ["gross"]}).status_code == 200
        state = client.get("/api/state").json()
        assert [f["key"] for f in state["fields"] if f["selected"]] == ["gross"]
        assert state["summary"]["fields"] == 1
        res = client.get("/api/export?format=xlsx")
        assert client.put("/api/selection", json={"keys": ["nope"]}).status_code == 400
    ws = openpyxl.load_workbook(io.BytesIO(res.content))["Invoices"]
    assert [c.value for c in ws[1]] == ["File", "Gross", "Status"]


def test_export_xlsx_download(tmp_path, server):
    xml = write_samples(tmp_path)[2]
    with client_for(resolve()) as client:
        upload(client, xml)
        wait_done(client)
        res = client.get("/api/export?format=xlsx")
        assert res.status_code == 200
        assert "invoices.xlsx" in res.headers["content-disposition"]
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        assert wb["Invoices"].max_row == 2
        assert wb["Invoices"]["A2"].value == "xrechnung.xml"
        assert "Line Items" in wb.sheetnames
        assert client.get("/api/export?format=txt").status_code == 400


def test_csv_export_with_table_is_zip(tmp_path, server):
    xml = write_samples(tmp_path)[2]
    with client_for(resolve()) as client:
        upload(client, xml)
        wait_done(client)
        client.put("/api/selection", json={"keys": ["gross", "line_items"]})
        res = client.get("/api/export?format=csv")
    assert res.headers["content-type"] == "application/zip"
    assert "invoices.zip" in res.headers["content-disposition"]
    names = zipfile.ZipFile(io.BytesIO(res.content)).namelist()
    assert sorted(names) == ["invoices-line_items.csv", "invoices.csv"]


def test_patch_cell_marks_edited_and_rechecks(tmp_path, server):
    xml = write_samples(tmp_path)[2]
    with client_for(resolve()) as client:
        file_id = upload(client, xml).json()["files"][0]["id"]
        wait_done(client)
        res = client.patch(f"/api/files/{file_id}/cells/gross", json={"value": "500.00"})
        assert res.status_code == 200
        assert client.patch(f"/api/files/{file_id}/cells/line_items", json={"value": "x"}).status_code == 400
        assert client.patch(f"/api/files/{file_id}/cells/nope", json={"value": "x"}).status_code == 404
        entry = client.get("/api/state").json()["files"][0]
    cell = entry["row"]["cells"]["gross"]
    assert cell["edited"] is True and cell["source"] == "manual"
    assert cell["value"] == "500.00" and cell["raw"] == "119.00" and cell["label"] == "Gross"
    assert entry["status"] == "error"
    assert any(c["name"] == "totals" and c["level"] == "error" for c in entry["row"]["checks"])


def test_template_mode_switch(tmp_path, server):
    a, _ = customer_pdfs(tmp_path)
    server.reply = discovery_reply()
    template = b"name: t\ncolumns:\n  - label: Order No.\n"
    with client_for(llm(server)) as client:
        upload(client, a)
        wait_done(client)
        res = client.post("/api/template", files={"file": ("t.yaml", template, "application/x-yaml")})
        assert res.status_code == 200
        state = wait_done(client)
        assert state["mode"] == "template"
        assert [f["key"] for f in state["fields"]] == ["order_no"]
        assert state["files"][0]["row"]["cells"]["order_no"]["value"] == "PO-1"
        bad = client.post("/api/template", files={"file": ("t.yaml", b"columns: 5", "application/x-yaml")})
        assert bad.status_code == 400
        assert client.delete("/api/template").status_code == 200
        state = wait_done(client)
    assert state["mode"] == "discover"
    assert [f["key"] for f in state["fields"]] == ["customer_no"]


def test_template_download_is_selection(tmp_path, server):
    xml = write_samples(tmp_path)[2]
    with client_for(resolve()) as client:
        upload(client, xml)
        wait_done(client)
        client.put("/api/selection", json={"keys": ["gross", "invoice_no"]})
        text = client.get("/api/template").text
    assert "Gross" in text and "Invoice No." in text and "IBAN" not in text
    assert 'Example values: \\"RE-2026-001\\"' in text or 'Example values: "RE-2026-001"' in text


def test_stale_generation_dropped(tmp_path, server):
    a, _ = customer_pdfs(tmp_path)
    gate = threading.Event()
    server.reply = discovery_reply(gate)
    template = b"name: t\ncolumns:\n  - label: Order No.\n"
    with client_for(llm(server)) as client:
        upload(client, a)
        wait_for(lambda: len(server.requests) > 0)
        client.post("/api/template", files={"file": ("t.yaml", template, "application/x-yaml")})
        gate.set()
        state = wait_done(client)
        time.sleep(0.3)
        state = wait_done(client)
    cells = state["files"][0]["row"]["cells"]
    assert "order_no" in cells and "kd_nr" not in cells
    assert [f["key"] for f in state["fields"]] == ["order_no"]


def test_unknown_id_404(server):
    with client_for(resolve()) as client:
        assert client.get("/api/files/nope/pages").status_code == 404
        assert client.get("/api/files/nope/pages/1.png").status_code == 404
        assert client.patch("/api/files/nope/cells/gross", json={"value": "1"}).status_code == 404


def test_upload_limit_413(tmp_path, server):
    paths = write_samples(tmp_path)
    with client_for(resolve(max_files=1)) as client:
        assert upload(client, paths[2], paths[1]).status_code == 413
        assert client.get("/api/state").json()["files"] == []
        assert upload(client, paths[2]).status_code == 200
        assert upload(client, paths[2]).status_code == 413
        assert client.delete("/api/files").status_code == 200
        state = client.get("/api/state").json()
        assert state["files"] == [] and state["fields"] == []


def test_index_served(server):
    with client_for(resolve()) as client:
        res = client.get("/")
        assert res.status_code == 200
        assert "text/html" in res.headers["content-type"]
        assert "<title>invoice-extract</title>" in res.text
        assert client.get("/static/app.js").status_code == 200
        assert client.get("/static/style.css").status_code == 200


def test_host_and_origin_guard(server):
    app = create_app(resolve(), allowed_hosts=["127.0.0.1", "localhost", "::1"])
    with TestClient(app, base_url="http://127.0.0.1:8000") as client:
        assert client.get("/api/state").status_code == 200
        assert client.get("/api/state", headers={"Host": "[::1]:8000"}).status_code == 200
        assert client.get("/api/state", headers={"Host": "evil.example"}).status_code == 400
        evil = {"Origin": "http://evil.example"}
        assert client.put("/api/selection", json={"keys": []}, headers=evil).status_code == 403
        same = {"Origin": "http://127.0.0.1:8000"}
        assert client.put("/api/selection", json={"keys": []}, headers=same).status_code == 200
    with client_for(resolve()) as client:
        evil = {"Origin": "http://evil.example"}
        assert client.put("/api/selection", json={"keys": []}, headers=evil).status_code == 403
