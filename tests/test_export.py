import csv
import datetime
import json
from decimal import Decimal

import openpyxl
import pytest

from invoice_extract import STANDARD_COLUMNS, Cell, Check, ConfigError, Row, export


def sample_row(supplier="ACME GmbH") -> Row:
    return Row(
        path="/in/inv-1.pdf",
        kind="pdf",
        cells={
            "file": Cell(value="inv-1.pdf", source="text"),
            "supplier": Cell(value=supplier, raw=supplier, page=1, source="text"),
            "invoice_no": Cell(value="RE-1", raw="RE-1", page=1, source="text", edited=True),
            "invoice_date": Cell(value=datetime.date(2026, 10, 1), raw="01.10.2026", page=1, source="text"),
            "gross": Cell(value=Decimal("1234.50"), raw="1.234,50", page=1, source="text",
                          checks=[Check("sum", "warning", "net + vat != gross")]),
        },
        checks=[Check("required", "warning", "Due date is missing")],
    )


def header(ws):
    return [c.value for c in ws[1]]


def test_xlsx_two_sheets_and_types(tmp_path):
    path = export([sample_row()], tmp_path / "out.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Invoices", "Details"]
    ws = wb["Invoices"]
    labels = header(ws)
    assert labels == [c.label for c in STANDARD_COLUMNS]
    values = dict(zip(labels, [c.value for c in ws[2]]))
    gross = ws.cell(2, labels.index("Gross") + 1)
    assert gross.value == 1234.5 and gross.number_format == "#,##0.00"
    assert isinstance(values["Invoice Date"], datetime.datetime)
    assert values["Status"] == "warning"
    assert values["File"] == "inv-1.pdf"
    assert ws.freeze_panes == "A2" and ws["A1"].font.bold

    details = wb["Details"]
    assert header(details) == ["File", "Column", "Value", "Raw text", "Source", "Page", "Edited", "Checks"]
    rows = [[c.value for c in r] for r in details.iter_rows(min_row=2)]
    by_column = {r[1]: r for r in rows}
    assert by_column["Gross"][7] == "warning: net + vat != gross"
    assert by_column["Invoice No."][6] in (True, "yes")
    assert by_column["(invoice)"][7] == "warning: Due date is missing"
    assert len(rows) == 6


def test_xlsx_formula_injection(tmp_path):
    path = export([sample_row('=HYPERLINK("x")')], tmp_path / "out.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Invoices"]
    cell = ws.cell(2, header(ws).index("Supplier") + 1)
    assert cell.data_type == "s"
    assert cell.quotePrefix or not str(cell.value).startswith("=")
    for row in wb["Details"].iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                assert c.data_type == "s" and c.quotePrefix


def test_csv_bom_and_injection(tmp_path):
    path = export([sample_row('=HYPERLINK("x")')], tmp_path / "out.csv")
    data = path.read_bytes()
    assert data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(data.decode("utf-8-sig").splitlines()))
    assert rows[0] == [c.label for c in STANDARD_COLUMNS]
    record = dict(zip(rows[0], rows[1]))
    assert record["Supplier"] == "'=HYPERLINK(\"x\")"
    assert record["Gross"] == "1234.50"
    assert record["Invoice Date"] == "2026-10-01"


def test_json_roundtrip(tmp_path):
    row = sample_row("Müller & Söhne")
    path = export([row], tmp_path / "out.json")
    text = path.read_text(encoding="utf-8")
    assert "Müller" in text
    data = json.loads(text)
    assert [c["key"] for c in data["columns"]] == [c.key for c in STANDARD_COLUMNS]
    assert data["rows"] == [row.to_dict()]


def test_error_row_exported_as_error(tmp_path):
    rows = [sample_row(), Row(path="/in/broken.pdf", error="UnsupportedFileError: not a PDF")]
    path = export(rows, tmp_path / "out.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Invoices"]
    values = dict(zip(header(ws), [c.value for c in ws[3]]))
    assert values["File"] == "broken.pdf"
    assert values["Status"] == "error"
    assert all(v is None for k, v in values.items() if k not in ("File", "Status"))
    errors = [r for r in wb["Details"].iter_rows(min_row=2, values_only=True) if r[1] == "(error)"]
    assert errors and errors[0][0] == "broken.pdf" and "not a PDF" in errors[0][7]

    path = export(rows, tmp_path / "out.csv")
    lines = list(csv.reader(path.read_text(encoding="utf-8-sig").splitlines()))
    assert lines[2][0] == "broken.pdf" and lines[2][-1] == "error"


def test_unknown_format(tmp_path):
    with pytest.raises(ConfigError):
        export([sample_row()], tmp_path / "out.txt")
    with pytest.raises(ConfigError):
        export([sample_row()], tmp_path / "out.xlsx", format="pdf")
    assert export([sample_row()], tmp_path / "out.dat", format="json").exists()


def test_export_survives_control_chars_and_non_scalar_values(tmp_path):
    # model output or PDF text can carry control characters, lists or objects; one bad cell must not stop the export
    from invoice_extract import Column
    cols = [Column("invoice_no", "Invoice No."), Column("gross", "Gross", type="amount"), Column("supplier", "Supplier")]
    row = Row(path="a.pdf", kind="pdf", cells={"invoice_no": Cell(value="abc\x02def"), "gross": Cell(value=[100]),
                                               "supplier": Cell(value={"name": "ACME"})})
    for fmt in ("xlsx", "csv", "json"):
        export([row], tmp_path / f"out.{fmt}", cols)
    ws = openpyxl.load_workbook(tmp_path / "out.xlsx")["Invoices"]
    assert [c.value for c in ws[2]] == ["abcdef", "[100]", '{"name": "ACME"}']


def test_xlsx_header_formula_is_text(tmp_path):
    from invoice_extract import Column
    label = '=HYPERLINK("http://e","c")'
    export([Row(path="a.pdf", cells={"x": Cell(value="1")})], tmp_path / "out.xlsx", [Column("x", label)])
    cell = openpyxl.load_workbook(tmp_path / "out.xlsx")["Invoices"]["A1"]
    assert cell.data_type == "s"
    assert cell.value == label
